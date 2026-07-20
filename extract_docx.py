import docx

doc = docx.Document('d:\\POSE\\XÂY DỰNG HỆ THỐNG QUẢN LÝ THƯƠNG MẠI ĐIỆN TỬ TÍCH HỢP CRM HỖ TRỢ HOẠT ĐỘNG BÁN HÀNG VÀ CHĂM SÓC KHÁCH HÀNG CHO DOANH NGHIỆP SME.docx')
with open('d:\\POSE\\srs_content.txt', 'w', encoding='utf-8') as f:
    for p in doc.paragraphs:
        f.write(p.text + '\n')
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                f.write(cell.text + ' | ')
            f.write('\n')
